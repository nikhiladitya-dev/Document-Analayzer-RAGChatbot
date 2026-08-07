from pathlib import Path

from langchain_core.documents import Document
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from app.core.logger import logger

class LoaderService:

    def load(self, file_path: str) -> list[Document]:
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"{file_path} not found.")

        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return self._load_pdf(path)

        elif suffix == ".docx":
            return self._load_docx(path)

        elif suffix == ".txt":
            return self._load_txt(path)

        raise ValueError(f"Unsupported file type: {suffix}")

    def _load_pdf(self, path: Path) -> list[Document]:

        logger.info(f"Loading PDF: {path.name}")

        documents = []

        pdf = fitz.open(path)

        for page_number, page in enumerate(pdf, start=1):

            text = page.get_text()

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": path.name,
                        "page": page_number,
                        "file_type": "pdf",
                    },
                )
            )

        pdf.close()

        return documents

    def _load_docx(self, path: Path) -> list[Document]:

        logger.info(f"Loading DOCX: {path.name}")

        doc = DocxDocument(path)

        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_type": "docx",
                },
            )
        ]

    def _load_txt(self, path: Path) -> list[Document]:

        logger.info(f"Loading TXT: {path.name}")

        text = path.read_text(
            encoding="utf-8"
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_type": "txt",
                },
            )
        ]